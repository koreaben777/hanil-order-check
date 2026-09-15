"""Local request-workbook -> sampled set plan -> schedule -> separate XLSX copy."""
from __future__ import annotations

import argparse
import json
import math
import random
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from hanilsf_plan import plan, to_xlsx
from hanilsf_schedule import schedule

__all__=['load_requests','sample_requests','main']


def load_requests(xlsx, sheets=None) -> list[dict]:
    """Read only numeric request rows; do not infer formulas or natural-language dates."""
    workbook=load_workbook(xlsx,read_only=True,data_only=True)
    requests=[]
    formula_book=None
    try:
        formula_book=load_workbook(xlsx,read_only=True,data_only=False)
        selected=workbook.sheetnames if sheets is None else ([sheets] if isinstance(sheets,str) else sheets)
        for name in selected:
            sheet=workbook[name]
            # Cached numeric values for requests, formula text only for the total boundary.
            total=next((r for r,values in enumerate(formula_book[name].iter_rows(min_row=9,max_col=8,values_only=True),9)
                        if str(values[7]).upper().startswith('=SUM(H9:')),None)
            if total is None: continue
            for row,values in enumerate(sheet.iter_rows(min_row=9,max_row=total-1,max_col=13,values_only=True),9):
                color=values[3]
                numbers=values[4:8]
                if not isinstance(color,str) or not re.fullmatch(r'[A-Za-z0-9]{4}',color.strip()): continue
                if not all(isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v) and v>0 for v in numbers): continue
                gsm,width,length,rolls=numbers
                if gsm!=int(gsm): raise ValueError('D-01 requires integer gsm for item code')
                gsm=int(gsm);color=color.strip().upper()
                # D-01: PP2 <=59gsm, PP3 >=60gsm; no ERP master lookup.
                request=dict(item='2PD'+('2' if gsm<=59 else '3')+f'{gsm:03d}'+color,
                             gsm=gsm,width=width,length=length,rolls=rolls,order_id=f'{name}:{row}',
                             partner_name=values[9] or '',end_user=values[10] or '',use=values[11] or '')
                due=values[12]
                if isinstance(due,(datetime,date)): request['due']=due.isoformat()
                elif isinstance(due,str) and due.strip().upper()=='ASAP': request['urgent']=True
                elif due is not None and str(due).strip(): request['note']=str(due)
                requests.append(request)
    finally:
        workbook.close()
        if formula_book is not None: formula_book.close()
    return requests


def sample_requests(pool,n,seed) -> list[dict]:
    return random.Random(seed).sample(pool,n)


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--requests',required=True)
    parser.add_argument('--n',type=int,default=30)
    parser.add_argument('--seed',type=int,default=1)
    parser.add_argument('--template')
    parser.add_argument('--out',required=True)
    parser.add_argument('--t0',default=datetime.combine(date.today()+timedelta(days=1),time()).isoformat())
    parser.add_argument('--state',required=True,help='color_code,gsm,agri e.g. WH1N,40,0')
    parser.add_argument('--time-limit',type=float,default=30)
    parser.add_argument('--past-due',choices=('ignore','keep'),default='ignore',
                        help='ignore sampled deadlines before t0 (default); keep for operational scheduling')
    args=parser.parse_args(argv)
    def report(**values): print(json.dumps(values,ensure_ascii=False,allow_nan=False))
    try:
        color,gsm,agri=[v.strip() for v in args.state.split(',')]
        if agri not in ('0','1'): raise ValueError('state agri must be 0 or 1')
        state=dict(color_code=color.upper(),gsm=float(gsm),agri=agri=='1')
        pool=load_requests(args.requests)
        rows=[dict(r) for r in sample_requests(pool,args.n,args.seed)]
        past_due_ignored=0
        if args.past_due=='ignore':
            start=datetime.fromisoformat(args.t0)
            for request in rows:
                if request.get('due'):
                    due=datetime.fromisoformat(request['due'])
                    if (due.tzinfo is None)!=(start.tzinfo is None):
                        raise ValueError('due and t0 must use consistent timezone awareness')
                    if due<start:
                        del request['due']
                        past_due_ignored+=1
        report(pool_rows=len(pool),sample_rows=len(rows),sample_items=len({r['item'] for r in rows}),
               seed=args.seed,past_due_ignored=past_due_ignored)
        result=plan(rows,time_limit=args.time_limit)
        statuses=sorted({g['status'] for g in result['groups']})
        report(stage=1,status=statuses,sets=sum(g['sets_used'] or 0 for g in result['groups']),
               total_length=sum(g['total_length'] or 0 for g in result['groups']))
        if not result['complete']:
            report(warnings=result['warnings'],output=None);return 1
        sequence=schedule(result,t0=args.t0,machine_state=state,time_limit=args.time_limit)
        report(stage=2,status=sequence['status'],jobs=len(sequence['jobs']),changeover_h=sequence['total_changeover_h'],
               cr_rolls=sequence['total_cr_rolls'],tardiness_h=sequence['total_tardiness_h'])
        if not sequence['complete']:
            report(warnings=sequence['warnings'],output=None);return 1
        out=to_xlsx(result,args.template or args.requests,args.out,schedule=sequence)
        report(warnings=list(dict.fromkeys(result['warnings']+sequence['warnings'])))
        report(output=str(Path(out).resolve()))
        return 0
    except (ValueError,OSError,KeyError) as exc:
        report(error=f'{type(exc).__name__}: {exc}',output=None)
        return 1


if __name__=='__main__':
    raise SystemExit(main())
