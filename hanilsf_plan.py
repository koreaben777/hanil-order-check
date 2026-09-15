"""Input-only slitting set planning. No ERP or external LLM connections.

Widths in mm, lengths in m. [미확인] K-01/K-06/K-08: equal-length sets,
no recutting. Set geometry is a data-backed assumption, not field approval.
"""
from __future__ import annotations

import math
import re
from collections import Counter, defaultdict
from copy import copy
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

__all__ = ["plan", "to_xlsx", "translate"]


def _integer(value, name, minimum=1):
    try:
        number=Decimal(str(value))
        if isinstance(value,bool) or not number.is_finite() or number!=number.to_integral_value() or number<minimum:
            raise ValueError()
        return int(number)
    except (InvalidOperation,TypeError,ValueError):
        raise ValueError(f"{name}: integer >= {minimum} required") from None


def _number(value, name, minimum=0, strict=False):
    try:
        n=float(value)
        if isinstance(value,bool) or not math.isfinite(n) or n<minimum or (strict and n==minimum): raise ValueError()
        return n
    except (ValueError,TypeError):
        raise ValueError(f"{name}: finite number required") from None


def _normalize(requests, warnings=None):
    orders,ids=[],set()
    for index,raw in enumerate(requests,1):
        if not isinstance(raw,dict): raise ValueError("requests must contain dicts")
        source=raw.get("order",raw)
        if not isinstance(source,dict): source=vars(source) if hasattr(source,"__dict__") else {}
        master=raw.get("master") or {}
        if not isinstance(master,dict): raise ValueError("master must be a dict")
        data={**source,**{k:v for k,v in raw.items() if k not in ("order","master")}}
        oid=data.get("order_id",f"O-{index:04d}")
        if not isinstance(oid,str) or not oid.strip() or oid in ids: raise ValueError("unique order_id required")
        ids.add(oid)
        item=str(data.get("item","")).strip().upper()
        if not re.fullmatch(r"[12][A-Z0-9_]+",item): raise ValueError("only machine 1/2 item codes supported")
        if len(item)!=11 and warnings is not None:
            warnings.append(f"item {item} is not 11 chars; color_code may be wrong")
        row=dict(order_id=oid,item=item,grade=data.get("grade","A"),
                 width=_integer(data.get("width"),"width"),length=_integer(data.get("length"),"length"),
                 rolls=_integer(raw.get("prod_rolls",data.get("rolls")),"rolls",0),
                 gsm=_number(data.get("gsm",master.get("gsm")),"gsm",strict=True))
        if row["grade"] is not None and not isinstance(row["grade"],str): raise ValueError("grade must be display text")
        for key in ("partner","partner_name","end_user","use","due","available_from","note","material","product_type"):
            value=data.get(key)
            row[key]=value.isoformat() if isinstance(value,date) else ("" if value is None else str(value))
        row["color_code"]=item[-4:]
        row["urgent"]=data.get("urgent",False)
        if not isinstance(row["urgent"],bool): raise ValueError("urgent must be boolean")
        row["mb_ratio"]=None if data.get("mb_ratio") is None else _number(data["mb_ratio"],"mb_ratio")
        if row["mb_ratio"] is not None and row["mb_ratio"]>1: raise ValueError("mb_ratio must be <= 1")
        orders.append(row)
    return orders


def _apply_constraints(orders,constraints):
    schema={"priority":{"type","order_id","rank"},"exclusive_roll":{"type","order_id"},
            "no_edge":{"type","order_id"},"trim":{"type","partner","mm"},
            "max_lanes":{"type","group","n"},"max_length":{"type","group","m"},
            "allow_extra":{"type","order_id","rolls"},"together":{"type","order_ids"},
            "eff_width":{"type","group","mm"},"unsupported":{"type","text"}}
    by_id={o["order_id"]:o for o in orders}
    groups={o["item"] for o in orders}
    partners={o["partner"] for o in orders if o["partner"]}
    applied,unsupported,seen=[],[],{}
    if not isinstance(constraints,(list,tuple)): raise ValueError("constraints must be a JSON list")
    for raw in constraints:
        if not isinstance(raw,dict) or not isinstance(raw.get("type"),str): raise ValueError("constraint type required")
        kind=raw["type"]
        if kind not in schema or set(raw)!=schema[kind]: raise ValueError("unknown constraint type or fields")
        c=dict(raw)
        for key in ("order_id","group","partner","text"):
            if key in c and (not isinstance(c[key],str) or not c[key].strip()): raise ValueError("invalid constraint reference")
        for key,valid in (("order_id",by_id),("group",groups),("partner",partners)):
            if key in c and c[key] not in valid: raise ValueError("unknown constraint reference")
        for key in ("rank","mm","n","m","rolls"):
            if key in c: c[key]=_integer(c[key],key,0 if key=="rolls" or (key=="mm" and kind=="trim") else 1)
        if kind=="together":
            ids=c["order_ids"]
            if not isinstance(ids,list) or len(ids)<2 or any(not isinstance(i,str) for i in ids) or len(set(ids))!=len(ids) or any(i not in by_id for i in ids):
                raise ValueError("together requires distinct valid order IDs")
            if len({by_id[i]["item"] for i in ids})!=1: raise ValueError("together requires same item")
            if len({by_id[i]["length"] for i in ids})!=1:
                unsupported.append(dict(type="unsupported",text="together: different lengths",order_ids=list(ids)))
                continue
            c["order_ids"]=list(ids)
        key=(kind,c.get("order_id",c.get("group",c.get("partner"))))
        if kind in ("priority","trim","max_length","max_lanes","allow_extra","eff_width"):
            if key in seen and seen[key]!=c: raise ValueError("conflicting scalar constraints")
            seen[key]=c
        target=unsupported if kind=="unsupported" else applied
        if c not in target: target.append(c)
    return applied,unsupported


def _group(orders,base_width_by_machine,eff_width_by_machine,eff_width_by_color,max_lanes,max_length,constraints):
    grouped=defaultdict(list)
    for o in orders:
        if o["rolls"]: grouped[o["item"]].append(o)
    groups=[]
    for item,rows in grouped.items():
        if len({o["gsm"] for o in rows})!=1: raise ValueError("same item must have same gsm")
        ratios={o["mb_ratio"] for o in rows}
        if len(ratios)>1: raise ValueError("same item must have consistent mb_ratio (including missing values)")
        ids={o["order_id"] for o in rows}
        cs=[c for c in constraints if c.get("order_id") in ids or c.get("group")==item
            or c.get("partner") in {o["partner"] for o in rows} or ids.intersection(c.get("order_ids",[]))]
        machine=item[0]
        if machine not in base_width_by_machine or machine not in eff_width_by_machine: raise ValueError("machine width missing")
        base,normal=base_width_by_machine[machine],eff_width_by_machine[machine]
        color=item[-4:]
        # [추정][미확인] W-03: agricultural exception; exact color overrides prefix.
        effective=eff_width_by_color.get(color,eff_width_by_color.get(color[:2],normal))
        source="color" if color in eff_width_by_color or color[:2] in eff_width_by_color else "machine"
        direct=next((c["mm"] for c in cs if c["type"]=="eff_width"),None)
        if direct is not None: effective,source=direct,"constraint"
        # W-01/W-02: trim only tightens; exception does not reduce safety trim requests.
        default_trim=(base-normal)/2
        trim=max([default_trim]+[c["mm"] for c in cs if c["type"]=="trim"])
        effective-=_integer(2*(trim-default_trim),"additional trim",0)
        if effective<=0 or effective>base: raise ValueError("effective width must be within production width")
        if trim>default_trim: source+="+trim"
        groups.append(dict(item=item,machine=machine,gsm=rows[0]["gsm"],base_width=base,
                           eff_width=effective,eff_width_source=source,
                           max_lanes=next((c["n"] for c in cs if c["type"]=="max_lanes"),max_lanes),
                           max_length=next((c["m"] for c in cs if c["type"]=="max_length"),max_length),
                           requests=rows,constraints=cs,mb_ratio=rows[0]["mb_ratio"]))
    return groups


def _patterns(rows,width,max_lanes,constraints):
    """All nonempty width multisets, not only maximal ones; filters are width-based."""
    widths=sorted({o["width"] for o in rows},reverse=True)
    by_id={o["order_id"]:o for o in rows}
    caps=Counter()
    for o in rows: caps[o["width"]]+=o["rolls"]+o["extra_cap"]
    exclusive={by_id[c["order_id"]]["width"] for c in constraints if c["type"]=="exclusive_roll" and c["order_id"] in by_id}
    edges={by_id[c["order_id"]]["width"] for c in constraints if c["type"]=="no_edge" and c["order_id"] in by_id}
    together=[{by_id[i]["width"] for i in c["order_ids"]} for c in constraints
              if c["type"]=="together" and all(i in by_id for i in c["order_ids"])]
    def valid(p):
        present={w for w,n in zip(widths,p) if n}
        if present & exclusive and len(present)>1: return False
        if any(bool(present & pair) and not pair<=present for pair in together): return False
        # [미확인] W-05: every protected-width lane must be internal, hence two unprotected lanes.
        if present & edges and (sum(p)<3 or sum(n for w,n in zip(widths,p) if w not in edges)<2): return False
        return True
    patterns=[]
    visited=0
    def enumerate_at(i,remaining,lanes,prefix):
        nonlocal visited
        if i==len(widths):
            if any(prefix):
                visited+=1
                if visited>50000: raise OverflowError
                if valid(prefix): patterns.append(tuple(prefix))
            return
        w=widths[i]
        for n in range(min(remaining//w,lanes,caps[w])+1):
            enumerate_at(i+1,remaining-n*w,lanes-n,prefix+[n])
    approximate=False
    try: enumerate_at(0,width,max_lanes,[])
    except OverflowError:
        # ponytail: >50,000 combinations: descending-width greedy seeds plus singleton
        # patterns only; report approximation, upgrade to column generation when needed.
        approximate=True
        patterns=[]
        for first in range(len(widths)):
            for n in range(1,min(width//widths[first],max_lanes,caps[widths[first]])+1):
                p=[0]*len(widths);p[first]=n
                if valid(p): patterns.append(tuple(p))
                remaining=width-n*widths[first]; lanes=max_lanes-n
                for i,w in enumerate(widths):
                    if i==first: continue
                    p[i]=min(remaining//w,lanes,caps[w]);remaining-=p[i]*w;lanes-=p[i]
                if valid(p): patterns.append(tuple(p))
    return widths,sorted(set(patterns)),edges,approximate


def _set_note(lanes,length,count):
    # W-05: only consecutive equal lanes may collapse; retain protected middle positions.
    from itertools import groupby
    runs=[(w,len(list(run))) for w,run in groupby(lanes)]
    pattern="+".join(str(w)+(f"*{n}" if n>1 else "") for w,n in runs)
    return pattern,f"({pattern})*{length} ×{count}"


def _solve_subproblem(rows,group,*,time_limit):
    """All stage-one CP-SAT code and the sole solver setting live here."""
    from ortools.sat.python import cp_model
    width=group["eff_width"]
    area=sum(o["width"]*o["rolls"] for o in rows)
    lower=(area+width-1)//width
    failure=dict(status="INFEASIBLE",sets=[],sets_used=None,total_length=None,
                 lower_bound_sets=lower,gap=None,approximate=False)
    if any(o["width"]>width or (group["max_length"] is not None and o["length"]>group["max_length"]) for o in rows): return failure
    widths,patterns,edges,approximate=_patterns(rows,width,group["max_lanes"],group["constraints"])
    failure["approximate"]=approximate
    if not patterns: return failure
    model=cp_model.CpModel()
    maximum=sum(o["rolls"]+o["extra_cap"] for o in rows)
    counts=[model.new_int_var(0,maximum,f"count_{i}") for i in range(len(patterns))]
    used=[model.new_bool_var(f"used_{i}") for i in range(len(patterns))]
    extra=[model.new_int_var(0,o["extra_cap"],f"extra_{i}") for i,o in enumerate(rows)]
    for x,u in zip(counts,used): model.add(x<=maximum*u); model.add(x>=u)
    # Same-width orders share lanes. Their exact totals/caps are conserved separately
    # on allocation, not duplicated as independent copies of the width demand.
    for k,w in enumerate(widths):
        model.add(sum(p[k]*x for p,x in zip(patterns,counts))==sum(o["rolls"]+e for o,e in zip(rows,extra) if o["width"]==w))
    extra_max=sum(o["extra_cap"] for o in rows)
    pattern_weight=extra_max+1
    set_weight=pattern_weight*(len(patterns)+1)
    model.minimize(set_weight*sum(counts)+pattern_weight*sum(used)+sum(extra))
    if model.validate(): raise ValueError("CP-SAT input/model integer range exceeded")
    solver=cp_model.CpSolver();solver.parameters.max_time_in_seconds=time_limit
    status=solver.solve(model)
    if status not in (cp_model.OPTIMAL,cp_model.FEASIBLE): return failure|dict(status=solver.status_name(status))
    remaining={o["order_id"]:o["rolls"] for o in rows}
    surplus={o["order_id"]:solver.value(e) for o,e in zip(rows,extra)}
    sets=[]
    for p,x in zip(patterns,counts):
        count=solver.value(x)
        if not count: continue
        lanes=[w for w,n in zip(widths,p) for _ in range(n)]
        if set(lanes)&edges:
            outer=[w for w in lanes if w not in edges]
            left,right=outer[0],outer[-1]
            lanes.remove(left);lanes.remove(right);lanes=[left]+lanes+[right]
        allocated,allocated_extra={},{}
        for w,n in zip(widths,p):
            quantity=n*count
            for pool,is_extra in ((remaining,False),(surplus,True)):
                for o in rows:
                    if o["width"]!=w: continue
                    oid=o["order_id"];take=min(quantity,pool[oid]);pool[oid]-=take;quantity-=take
                    if take:
                        allocated[oid]=allocated.get(oid,0)+take
                        if is_extra: allocated_extra[oid]=take
            if quantity: raise ValueError("width allocation imbalance")
        pattern,note=_set_note(lanes,rows[0]["length"],count)
        sets.append(dict(pattern=pattern,lanes=lanes,width_sum=sum(lanes),loss_mm=width-sum(lanes),
                         length=rows[0]["length"],count=count,rolls=allocated,extra=allocated_extra,note=note))
    total=sum(s["count"] for s in sets)
    # OPTIMAL applies to enumerated candidates only if fallback was not used.
    bound=total if status==cp_model.OPTIMAL and not approximate else lower
    return dict(status=solver.status_name(status),sets=sets,sets_used=total,
                total_length=total*rows[0]["length"],lower_bound_sets=bound,
                gap=(total-bound)/total,approximate=approximate)


def _validate_result(result):
    def require(ok):
        if not ok: raise ValueError("set plan invariant violated")
    orders={o["order_id"]:o for o in result["orders"]}
    counts,extras=Counter(),Counter()
    solved=set()
    for g in result["groups"]:
        if g["status"] not in ("OPTIMAL","FEASIBLE"):
            require(not g["sets"] and g["total_length"] is None);continue
        solved.add(g["item"])
        require(g["sets_used"]==sum(s["count"] for s in g["sets"]))
        require(g["total_length"]==sum(s["count"]*s["length"] for s in g["sets"]))
        require(0<=g["lower_bound_sets"]<=g["sets_used"])
        require(g['mr_length_m']==g['total_length'])
        require(math.isclose(g['mr_weight_kg'],sum(s['mr_weight_kg'] for s in g['sets'])))
        if g['production_hours'] is not None:
            require(all(s['production_hours'] is not None for s in g['sets']))
            require(math.isclose(g['production_hours'],sum(s['production_hours'] for s in g['sets'])))
        if g['mb_kg'] is not None:
            require(math.isclose(g['mb_kg'],sum(s['mb_kg'] for s in g['sets'])))
        cs=result['constraints_applied']
        edges={orders[c['order_id']]['width'] for c in cs if c['type']=='no_edge' and orders[c['order_id']]['item']==g['item']}
        exclusive={orders[c['order_id']]['width'] for c in cs if c['type']=='exclusive_roll' and orders[c['order_id']]['item']==g['item']}
        for s in g["sets"]:
            require(isinstance(s["count"],int) and s["count"]>0)
            require(0<len(s["lanes"])<=g["max_lanes"] and all(isinstance(w,int) and w>0 for w in s["lanes"]))
            require(sum(s["lanes"])==s["width_sum"]<=g["eff_width"])
            require(s["loss_mm"]==g["eff_width"]-s["width_sum"])
            require(g["max_length"] is None or s["length"]<=g["max_length"])
            require(s['pattern']==_set_note(s['lanes'],s['length'],s['count'])[0])
            require(s['mr_length_m']==s['length']*s['count'])
            present=set(s['lanes'])
            require(not present.intersection(edges) or (len(s['lanes'])>=3 and s['lanes'][0] not in edges and s['lanes'][-1] not in edges))
            require(not present.intersection(exclusive) or len(present)==1)
            for c in cs:
                if c['type']=='together' and orders[c['order_ids'][0]]['item']==g['item'] and orders[c['order_ids'][0]]['length']==s['length']:
                    pair={orders[i]['width'] for i in c['order_ids']}
                    require(not present.intersection(pair) or pair<=present)
            width_counts=Counter()
            for oid,n in s["rolls"].items():
                require(oid in orders and isinstance(n,int) and n>0)
                o=orders[oid];require(o["item"]==g["item"] and o["length"]==s["length"])
                counts[oid]+=n;width_counts[o["width"]]+=n
            require(width_counts==Counter({w:n*s["count"] for w,n in Counter(s["lanes"]).items()}))
            for oid,n in s["extra"].items():
                require(isinstance(n,int) and 0<n<=s["rolls"].get(oid,0));extras[oid]+=n
    for oid,o in orders.items():
        require(o["placed"]==counts[oid] and o["extra"]==extras[oid]<=o["extra_cap"])
        require(o["suggest_upsell"]==(extras[oid]>0))
        if o["item"] in solved or result["complete"]: require(counts[oid]==o["rolls"]+extras[oid])
    require(math.isclose(result["total_kg"],sum(o["gsm"]*o["width"]*o["length"]*counts[oid]/1e6 for oid,o in orders.items())))


# W-01/W-02 confirmed widths; [추정][미확인] W-03 agricultural 3500 exception.
# [추정] F-02/F-03: 3.654, 27 and gsm k are caller-replaceable, not standards.
def plan(requests, *, base_width_by_machine={"1":3400,"2":3600},
         eff_width_by_machine={"1":3200,"2":3400},eff_width_by_color={"UV":3500,"UB":3500},
         max_lanes=30,max_length=None,constraints=(),time_limit=30,mr_width_factor=3.654,
         time_divisors=(27,{**{g:26 for g in range(30,91)},15:21,18:21,100:21,140:19})) -> dict:
    """Plan equal-length sets; extras forbidden unless allow_extra grants an order cap.

    Unknown gsm k uses the nearest supplied gsm (ties: smaller gsm), with a warning.
    Approximate pattern fallback is marked explicitly, even if restricted IP is OPTIMAL.
    """
    maps=[]
    for mapping in (base_width_by_machine,eff_width_by_machine,eff_width_by_color):
        if not isinstance(mapping,dict): raise ValueError("width mappings must be dicts")
        maps.append({str(k):_integer(v,"width") for k,v in mapping.items()})
    max_lanes=_integer(max_lanes,"max_lanes")
    if max_length is not None: max_length=_integer(max_length,"max_length")
    seconds=_number(time_limit,"time_limit",strict=True)
    factor=_number(mr_width_factor,"mr_width_factor",strict=True)
    if not isinstance(time_divisors,(list,tuple)) or len(time_divisors)!=2 or not isinstance(time_divisors[1],dict): raise ValueError("time_divisors requires (divisor, gsm:k dict)")
    divisor=_number(time_divisors[0],"divisor",strict=True)
    ks={_number(k,"gsm",strict=True):_number(v,"k",strict=True) for k,v in time_divisors[1].items()}
    if not ks: raise ValueError("time_divisors gsm:k table must not be empty")
    warnings=[]
    orders=_normalize(requests,warnings)
    applied,unsupported=_apply_constraints(orders,constraints)
    caps={c["order_id"]:c["rolls"] for c in applied if c["type"]=="allow_extra"}
    for o in orders: o["extra_cap"]=caps.get(o["order_id"],0)
    groups=_group(orders,*maps,max_lanes,max_length,applied)
    warnings.extend(["[미확인] K-01/K-06/K-08: equal-length set structure; no recutting", "[추정] F-02/F-03: production coefficients"])
    if unsupported: warnings.append("Unsupported constraints were not applied")
    counts,extras=Counter(),Counter()
    for g in groups:
        # [추정] F-03: use only original table keys, never previously inferred gsm values.
        nearest=min(ks,key=lambda gsm:(abs(gsm-g["gsm"]),gsm))
        k=ks[nearest]
        if g["gsm"] not in ks:
            warnings.append(f"[추정] F-03: gsm {g['gsm']:g} uses k of {nearest:g}")
        parts=defaultdict(list)
        for o in g["requests"]: parts[o["length"]].append(o)
        results=[_solve_subproblem(rows,g,time_limit=seconds) for rows in parts.values()]
        complete=all(r["status"] in ("OPTIMAL","FEASIBLE") for r in results)
        status="OPTIMAL" if all(r["status"]=="OPTIMAL" for r in results) else ("FEASIBLE" if complete else ("INFEASIBLE" if any(r["status"]=="INFEASIBLE" for r in results) else "UNKNOWN"))
        sets=[s for r in results for s in r["sets"]] if complete else []
        ranks={c["order_id"]:c["rank"] for c in g["constraints"] if c["type"]=="priority"}
        sets.sort(key=lambda s:(min((ranks.get(i,math.inf) for i in s["rolls"]),default=math.inf),s["length"],s["pattern"]))
        for i,s in enumerate(sets,1):
            s["set_no"]=i
            s["mr_length_m"]=s["length"]*s["count"]
            s["mr_weight_kg"]=s["mr_length_m"]*factor*g["gsm"]/1000
            s["production_hours"]=s["mr_weight_kg"]/divisor/k
            s["mb_kg"]=s["mr_weight_kg"]*g["mb_ratio"] if g["mb_ratio"] is not None else None
            counts.update(s["rolls"]);extras.update(s["extra"])
        total=sum(s["mr_length_m"] for s in sets) if complete else None
        used=sum(s["count"] for s in sets) if complete else None
        lower=sum(r["lower_bound_sets"] for r in results)
        g.update(status=status,sets=sets,total_length=total,sets_used=used,lower_bound_sets=lower,
                 gap=(used-lower)/used if complete else None,approximate=any(r["approximate"] for r in results),
                 mr_length_m=total,mr_weight_kg=sum(s["mr_weight_kg"] for s in sets) if complete else None,
                 production_hours=sum(s["production_hours"] for s in sets) if complete else None,
                 mb_kg=sum(s["mb_kg"] for s in sets) if complete and g["mb_ratio"] is not None else None,
                 coefficients_estimated=True)
        if g["approximate"]: warnings.append("Approximate greedy pattern fallback; restricted-model status only")
        g.pop("requests");g.pop("constraints")
    result=dict(groups=groups,orders=[o|dict(placed=counts[o["order_id"]],extra=extras[o["order_id"]],
                suggest_upsell=extras[o["order_id"]]>0,unplaced=max(0,o["rolls"]+extras[o["order_id"]]-counts[o["order_id"]])) for o in orders],
                complete=all(g["status"] in ("OPTIMAL","FEASIBLE") for g in groups),
                total_kg=sum(o["gsm"]*o["width"]*o["length"]*counts[o["order_id"]]/1e6 for o in orders),
                constraints_applied=applied,constraints_unsupported=unsupported,coefficients_estimated=True,
                warnings=list(dict.fromkeys(warnings)))
    _validate_result(result)
    return result


def translate(text,orders) -> list[dict]:
    raise NotImplementedError("Pass manually reviewed JSON through constraints")


def _insert_rows(sheet, at, amount):
    """openpyxl은 행 삽입 시 수식·병합 참조를 갱신하지 않으므로 명시적으로 이동."""
    from openpyxl.formula import Tokenizer
    from openpyxl.worksheet.cell_range import CellRange

    merges = [CellRange(str(m)) for m in sheet.merged_cells]
    for m in list(sheet.merged_cells):
        sheet.unmerge_cells(str(m))
    dimensions = {r: copy(d) for r,d in sheet.row_dimensions.items()}
    sheet.insert_rows(at, amount)
    sheet.row_dimensions.clear()
    for r,d in dimensions.items():
        target = r+amount if r>=at else r
        d.index = target
        sheet.row_dimensions[target] = d
    for m in merges:
        if m.min_row >= at:
            m.shift(row_shift=amount)
        elif m.max_row >= at:
            m.max_row += amount
        sheet.merge_cells(str(m))
    for row in sheet:
        for cell in row:
            if cell.data_type != "f":
                continue
            # 단일 셀 색상 집계도 블록 확장에 포함(단순 =I행은 SUM 범위로 변환).
            formula = cell.value
            if cell.row >= at+amount:
                single = re.fullmatch(r"=\$?I\$?(\d+)|=SUM\(\$?I\$?(\d+)\)", formula, re.IGNORECASE)
                if single and int(single.group(1) or single.group(2)) == at-1:
                    formula = f"=SUM(I{at-1}:I{at-1})"
            tokens = Tokenizer(formula).items
            for t in tokens:
                if t.type != "OPERAND" or t.subtype != "RANGE":
                    continue
                prefix, sep, ref = t.value.rpartition("!")
                if sep and prefix.strip("'").replace("''", "'") != sheet.title:
                    continue
                match = re.fullmatch(r"(\$?[A-Z]{1,3}\$?)(\d+)(?::(\$?[A-Z]{1,3}\$?)(\d+))?", ref)
                if not match:
                    continue
                c1,r1,c2,r2 = match.groups()
                first = int(r1)
                last = int(r2) if r2 else None
                new_first = first+amount if first>=at else first
                if last is not None:
                    # 블록 끝 바로 앞 삽입은 그 블록을 합산하는 범위에도 포함한다.
                    new_last = last+amount if last>=at or (last==at-1 and first<at) else last
                    ref = f"{c1}{new_first}:{c2}{new_last}"
                else:
                    ref = f"{c1}{new_first}"
                t.value = (prefix+sep if sep else "")+ref
            cell.value = "="+"".join(t.value for t in tokens)


def to_xlsx(result, template, out_path, sheet_name=None, schedule=None) -> Path:
    """PP 양식의 첫 구조 일치 시트를 복제·비우고 사본에만 쓴다. 기존 출력도 덮어쓰지 않는다.

    color_code(기본 품목코드 뒤 4자리)로 블록을 찾는다. 불명확하면 실패한다.
    수식은 유지하고 재계산을 요청한다; openpyxl 자체는 Excel 수식을 계산하지 않는다.
    """
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.workbook.properties import CalcProperties

    _validate_result(result)
    if not result["complete"]:
        raise ValueError("미배치 그룹이 있는 계획은 생산의뢰서로 출력할 수 없습니다")
    template, out_path = Path(template), Path(out_path)
    if template.resolve() == out_path.resolve() or out_path.exists():
        raise ValueError("원본 또는 기존 출력 파일을 덮어쓸 수 없습니다")
    if schedule is not None:
        if not schedule.get("complete") or schedule.get("status") not in ("OPTIMAL","FEASIBLE"):
            raise ValueError("incomplete schedule cannot be exported")
        expected={(g["item"],s["set_no"]):(s["pattern"],s["length"],s["count"],s["rolls"])
                  for g in result["groups"] for s in g["sets"]}
        actual={(j["item"],j["set_no"]):(j["pattern"],j["length"],j["count"],j["rolls"]) for j in schedule["jobs"]}
        if expected!=actual or len(actual)!=len(schedule["jobs"]): raise ValueError("schedule does not match plan sets")
    workbook = load_workbook(template)
    try:
        source = next((s for s in workbook if s["H8"].value == "roll" and s["I8"].value == "kgs"
                       and any(str(s.cell(r,8).value).upper().startswith("=SUM(H9:") for r in range(9,s.max_row+1))), None)
        if source is None:
            raise ValueError("PP 템플릿 시트 구조를 찾을 수 없습니다")
        name = sheet_name or date.today().strftime("%Y%m%d(%y.%m.%d)")
        if not name or len(name)>31 or re.search(r"[\\/*?:\[\]]", name) or name in workbook.sheetnames:
            raise ValueError("새롭고 유효한 sheet_name이 필요합니다")
        reserved={"세트표","스케줄표"}
        if name in reserved or reserved.intersection(workbook.sheetnames):
            raise ValueError("세트표/스케줄표 names are reserved; existing sheets are never overwritten")
        sheet = workbook.copy_worksheet(source)
        sheet.title = name
        total = next(r for r in range(9,sheet.max_row+1) if str(sheet.cell(r,8).value).upper().startswith("=SUM(H9:"))
        # 구조만 복제한다. 기존 회차의 주문·현재고·거래선·비고를 새 시트에 남기지 않는다.
        for row in sheet.iter_rows(min_row=9,max_row=total-1):
            for cell in row:
                if not isinstance(cell,MergedCell) and cell.column != 4:
                    if cell.column != 9 or cell.data_type != "f":
                        cell.value = None
                    cell.comment = None
                    cell.hyperlink = None
        by_id = {o["order_id"]: o for o in result["orders"]}
        entries = defaultdict(list)
        for g in result["groups"]:
            for block in g["sets"]:
                for oid,count in block["rolls"].items():
                    extra=block["extra"].get(oid,0)
                    if count-extra:
                        entries[by_id[oid]["color_code"]].append((g,block,by_id[oid],False,count-extra))
                    if extra:
                        entries[by_id[oid]["color_code"]].append((g,block,by_id[oid],True,extra))
                entries[g["item"][-4:]].append(None)
        # 실제 블록 경계는 D열의 연속 색상 라벨; 빈 구분 행은 이전 블록에 속한다.
        blocks = []
        for r in range(9,total):
            label = sheet.cell(r,4).value
            if label is not None and (not blocks or blocks[-1][2] != str(label)):
                blocks.append([r,total,str(label)])
        for a,b in zip(blocks,blocks[1:]):
            a[1] = b[0]
        jobs = []
        for color, rows in entries.items():
            candidates = [b for b in blocks if b[2] == color]
            if not candidates:
                candidates = [b for b in blocks if re.search(rf"(?<![A-Z0-9]){re.escape(color)}(?![A-Z0-9])",b[2])]
            if len(candidates)>1:
                raise ValueError("색상 블록을 유일하게 찾을 수 없습니다; color_code를 확인하세요")
            while rows and rows[-1] is None:
                rows.pop()
            if not candidates:
                # New colors belong immediately before totals; preserve last-block formatting.
                start=total
                style=[copy(sheet.cell(total-1,c)._style) for c in range(1,sheet.max_column+1)]
                dimension=copy(sheet.row_dimensions[total-1])
                _insert_rows(sheet,total,len(rows))
                for r in range(start,start+len(rows)):
                    for c,st in enumerate(style,1): sheet.cell(r,c)._style=copy(st)
                    dim=copy(dimension);dim.index=r;sheet.row_dimensions[r]=dim
                    sheet.cell(r,4,color)
                total+=len(rows)
                candidates=[[start,total,color]]
                warning=f"template has no block for {color}; appended"
                if warning not in result["warnings"]: result["warnings"].append(warning)
            jobs.append((*candidates[0][:2],rows))
        # 아래 블록부터 채워야 위 블록 행 삽입에도 이미 쓴 데이터와 수식이 함께 이동한다.
        for start,end,rows in sorted(jobs,reverse=True):
            if len(rows)>end-start:
                amount = len(rows)-(end-start)
                style = [copy(sheet.cell(end-1,c)._style) for c in range(1,sheet.max_column+1)]
                _insert_rows(sheet,end,amount)
                for r in range(end,end+amount):
                    for c,st in enumerate(style,1):
                        sheet.cell(r,c)._style = copy(st)
                    sheet.cell(r,4,sheet.cell(start,4).value)
                total += amount
            for r,entry in enumerate(rows,start):
                if entry is None:
                    for c in range(1,20):
                        sheet.cell(r,c).value = None
                    continue
                g,block,o,extra,count = entry
                note=f'{block["note"]}세트 #세트{block["set_no"]}'
                if extra: note+=f" · 초과 {count}롤 · 추가구매 권유"
                values = {1:o["material"],2:int(o["item"][0]),3:o["product_type"],4:o["color_code"],
                          5:o["gsm"],6:o["width"],7:o["length"],8:count,
                          10:o["partner_name"],11:o["end_user"],12:o["use"],13:o["due"],19:note}
                for c,v in values.items():
                    cell = sheet.cell(r,c)
                    cell.value = v
                    if isinstance(v,str):
                        cell.data_type = "s"  # 거래선/비고 '=...'는 수식으로 실행하지 않는다.
                sheet.cell(r,9,f"=E{r}/1000*F{r}/1000*G{r}*H{r}")
        sheet["A2"] = f"생산의뢰 NO.: {name} · [미확인] 세트 구조(K-01) · 계수 추정(F-02/F-03)"
        sheet.print_area = f"A1:S{sheet.max_row}"
        sets_sheet=workbook.create_sheet("세트표")
        sets_sheet.append(["세트번호","품목코드","패턴","폭 합","로스(mm)","길이(m)","세트 수",
                           "주문별 롤수","M/R 길이","M/R 중량","생산시간(추정)","MB 소요량"])
        import json
        for g in result["groups"]:
            for block in g["sets"]:
                sets_sheet.append([block["set_no"],g["item"],block["pattern"],block["width_sum"],block["loss_mm"],
                                   block["length"],block["count"],json.dumps(block["rolls"],ensure_ascii=False),
                                   block["mr_length_m"],block["mr_weight_kg"],block["production_hours"],block["mb_kg"]])
        output_sheets=[sets_sheet]
        if schedule is not None:
            sched=workbook.create_sheet("스케줄표")
            sched.append(["순번","품목코드","색상","gsm","패턴","길이","세트 수","롤수","M/R 길이",
                          "M/R 중량","생산시간","시작","종료","전환(h)","CR 롤","납기","지연(h)"])
            changes={tr["to"]:tr for tr in schedule["transitions"]}
            for j in schedule["jobs"]:
                tr=changes[j["job_id"]]
                sched.append([j["sequence"],j["item"],j["color_code"],j["gsm"],j["pattern"],j["length"],j["count"],
                              sum(j["rolls"].values()),j["mr_length_m"],j["mr_weight_kg"],j["production_hours"],
                              j["start"],j["end"],tr["delta_h"],tr["cr_rolls"],j["due"],j["tardiness_h"]])
            output_sheets.append(sched)
        for output in output_sheets:
            output.freeze_panes="A2"
            for row in output:
                for cell in row:
                    if isinstance(cell.value,str): cell.data_type="s"
        workbook.calculation = CalcProperties(fullCalcOnLoad=True)
        # x 모드로 저장하여 경로 검사 이후에도 기존 파일을 덮어쓰지 않는다.
        with out_path.open("xb") as stream:
            workbook.save(stream)
    finally:
        workbook.close()
    return out_path


