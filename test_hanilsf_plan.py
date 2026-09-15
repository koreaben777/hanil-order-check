"""DB-free, synthetic self-tests: python test_hanilsf_plan.py [--core]."""
import copy
import json
import math
import plistlib
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path
from types import SimpleNamespace

import hanilsf_plan as hp
from openpyxl import Workbook, load_workbook


def order(oid="a", width=1030, length=1000, rolls=18, **kw):
    return dict(order_id=oid, item="2TESTWH1N", grade="A", width=width,
                length=length, rolls=rolls, gsm=40) | kw


def fails(fn, exc=ValueError):
    try:
        fn()
    except exc:
        return
    raise AssertionError(f"expected {exc.__name__}")


def verify(r):
    assert r["complete"]
    counts, extras = Counter(), Counter()
    orders = {o["order_id"]: o for o in r["orders"]}
    for g in r["groups"]:
        assert g["status"] in ("OPTIMAL", "FEASIBLE")
        assert g["sets_used"] == sum(s["count"] for s in g["sets"])
        assert g["total_length"] == sum(s["length"]*s["count"] for s in g["sets"])
        assert 0 <= g["lower_bound_sets"] <= g["sets_used"]
        for s in g["sets"]:
            assert sum(s["lanes"]) == s["width_sum"] <= g["eff_width"]
            assert 1 <= len(s["lanes"]) <= g["max_lanes"]
            assert sum(s["rolls"].values()) == len(s["lanes"])*s["count"]
            widths = Counter()
            for oid, n in s["rolls"].items():
                assert orders[oid]["item"] == g["item"] and orders[oid]["length"] == s["length"]
                counts[oid] += n
                widths[orders[oid]["width"]] += n
            assert widths == Counter({w: n*s["count"] for w,n in Counter(s["lanes"]).items()})
            extras.update(s["extra"])
    for o in r["orders"]:
        oid=o["order_id"]
        assert counts[oid] == o["placed"] == o["rolls"]+o["extra"]
        assert extras[oid] == o["extra"] <= o["extra_cap"]
        assert o["suggest_upsell"] == (o["extra"]>0)
    assert math.isclose(r["total_kg"],sum(o["gsm"]*o["width"]*o["length"]*o["placed"]/1e6 for o in r["orders"]))
    assert json.dumps(r,allow_nan=False)
    hp._validate_result(r)
    return r


def test_a_known_sets_and_grade_passthrough():
    r=verify(hp.plan([order(),order("b",530,grade="B")]))
    g=r["groups"][0]
    assert len(r["groups"])==1 and r["orders"][1]["grade"]=="B"
    assert g["status"]=="OPTIMAL" and g["sets_used"]==9 and g["total_length"]==9000
    assert all(o["extra"]==0 for o in r["orders"])
    print("  (a) 1030x18 + 530x18: OPTIMAL, sets=9, length=9000, extra=0")
    r=verify(hp.plan([order(width=1300,rolls=2),order("b",600,rolls=1)],eff_width_by_machine={"2":3200}))
    assert r["groups"][0]["sets_used"]==1 and r["groups"][0]["sets"][0]["width_sum"]==3200
    r=verify(hp.plan([order(rolls=7)]))
    assert r["groups"][0]["sets_used"]==3  # non-maximal patterns required
    r=verify(hp.plan([order(rolls=1),order("b",length=500,rolls=1)]))
    assert r["groups"][0]["total_length"]==1500


def test_b_constraints():
    rows=[order(width=1800,rolls=1,partner="SYNTH"),order("b",1500,rolls=1)]
    def solve(cs=(), **kw): return verify(hp.plan(rows,constraints=cs,**kw))["groups"][0]
    assert solve()["sets_used"]==1
    assert solve([dict(type="exclusive_roll",order_id="a")])["sets_used"]==2
    assert solve([dict(type="max_lanes",group="2TESTWH1N",n=1)])["sets_used"]==2
    assert solve([dict(type="trim",partner="SYNTH",mm=200)])["eff_width"]==3200
    assert solve([dict(type="eff_width",group="2TESTWH1N",mm=3200)])["sets_used"]==2
    g=solve([dict(type="priority",order_id="b",rank=1)],max_lanes=1)
    assert "b" in g["sets"][0]["rolls"]
    r=hp.plan(rows,constraints=[dict(type="max_length",group="2TESTWH1N",m=999)])
    assert not r["complete"] and r["groups"][0]["status"]=="INFEASIBLE"
    rows2=[order(width=1000,rolls=1),order("b",600,rolls=2)]
    r=verify(hp.plan(rows2,constraints=[dict(type="no_edge",order_id="a")]))
    assert r["groups"][0]["sets"][0]["lanes"]==[600,1000,600]
    assert r["groups"][0]["sets"][0]["pattern"]=='600+1000+600'  # note must preserve protected middle lane
    r=hp.plan([order(width=1000,rolls=3)],constraints=[dict(type="no_edge",order_id="a")])
    assert not r["complete"]  # all protected widths cannot occupy outer lanes
    cs=[dict(type="together",order_ids=["a","b"])]
    r=verify(hp.plan(rows2,constraints=cs))
    assert all(1000 in s["lanes"] and 600 in s["lanes"] for s in r["groups"][0]["sets"])
    r=verify(hp.plan([rows2[0],rows2[1]|dict(length=500)],constraints=cs))
    assert len(r["constraints_unsupported"])==1 and not r["constraints_applied"]
    print("  (b) priority/exclusive_roll/no_edge/trim/max_lanes/max_length/together/eff_width/unsupported: passed")


def test_c_exact_and_allow_extra():
    rows=[order(width=1600,rolls=3)]
    r=verify(hp.plan(rows))
    assert len(r["groups"][0]["sets"])==2 and r["orders"][0]["extra"]==0
    r=verify(hp.plan(rows,constraints=[dict(type="allow_extra",order_id="a",rolls=1)]))
    assert r["groups"][0]["sets_used"]==2 and len(r["groups"][0]["sets"])==1
    assert r["orders"][0]["extra"]==1 and r["orders"][0]["suggest_upsell"]
    r=verify(hp.plan([order(rolls=2),order("b",rolls=1)],constraints=[dict(type="allow_extra",order_id="b",rolls=1)]))
    assert all(o["extra"]==0 for o in r["orders"])  # same width pooled, no unnecessary extras
    print("  (c) exact default=0 extra; allow_extra cap=1 => extra=1, fewer patterns")


def test_d_agricultural_width():
    rows=[order(width=1800,rolls=1,item="2TESTUV2N"),order("b",1700,rolls=1,item="2TESTUV2N")]
    g=verify(hp.plan(rows))["groups"][0]
    assert g["eff_width"]==3500 and g["eff_width_source"]=="color" and g["sets_used"]==1
    g=verify(hp.plan(rows,eff_width_by_color={}))["groups"][0]
    assert g["eff_width"]==3400 and g["sets_used"]==2
    assert verify(hp.plan([order(item="1TESTWH1N")]))["groups"][0]["eff_width"]==3200
    print("  (d) UV width=3500: sets=1; exception disabled: sets=2; machine1 width=3200")


def test_validation_metrics_and_failure():
    raw=dict(order=SimpleNamespace(**order(rolls=99)),prod_rolls=2,master=dict(gsm=40))
    before=copy.deepcopy(raw)
    r=verify(hp.plan([raw]))
    assert raw==before and r["orders"][0]["rolls"]==2
    r=verify(hp.plan([order(width=3400,rolls=1,mb_ratio=.05)]))
    g=r["groups"][0]
    assert math.isclose(g["mr_weight_kg"],146.16)
    assert math.isclose(g["production_hours"],146.16/27/26)
    assert math.isclose(g["mb_kg"],7.308) and r["coefficients_estimated"]
    assert hp.plan([])["complete"]
    for patch in [dict(width=0),dict(length=True),dict(rolls=-1),dict(gsm=None),dict(gsm=float('nan')),dict(item="3TEST")]:
        fails(lambda patch=patch: hp.plan([order()|patch]))
    for cs in [[dict(type="no_overproduce",order_id="a")],[dict(type="allow_extra",order_id="missing",rolls=1)],
               [dict(type="priority",order_id="a",rank=1,unknown=True)]]:
        fails(lambda cs=cs: hp.plan([order()],constraints=cs))
    fails(lambda: hp.plan([order(),order()]))
    fails(lambda: hp.plan([order()],time_limit=0))
    fails(lambda: hp.translate("manual",[]),NotImplementedError)
    r=hp.plan([order()],time_limit=1e-9)
    assert not r["complete"] and r["groups"][0]["status"]=="UNKNOWN"
    fails(lambda: hp.to_xlsx(r,"not-read.xlsx","not-written.xlsx"))
    bad=copy.deepcopy(verify(hp.plan([order()])))
    bad["groups"][0]["sets"][0]["count"]+=1
    fails(lambda: hp._validate_result(bad))
    bad=copy.deepcopy(verify(hp.plan([order(width=1000,rolls=1),order('b',600,rolls=2)],
                                    constraints=[dict(type='no_edge',order_id='a')])))
    bad['groups'][0]['sets'][0]['lanes']=[1000,600,600]
    fails(lambda: hp._validate_result(bad))
    bad=copy.deepcopy(verify(hp.plan([order()])))
    bad['groups'][0]['mr_weight_kg']+=1
    fails(lambda: hp._validate_result(bad))
    assert 'hanilsf_optimizer' not in sys.modules and 'hhhs_db_manager' not in sys.modules


def test_nearest_gsm_and_item_warning():
    r=verify(hp.plan([order(width=3400,rolls=1,gsm=20)]))
    hours=r['groups'][0]['production_hours']
    assert isinstance(hours,(int,float)) and math.isclose(hours,73.08/27/21)
    assert '[추정] F-03: gsm 20 uses k of 18' in r['warnings']
    assert 'item 2TESTWH1N is not 11 chars; color_code may be wrong' in r['warnings']
    assert 'item 2TESTWH1N is not 11 chars; color_code may be wrong' in hp.plan([order(rolls=0)])['warnings']
    r2=verify(hp.plan([order(width=3400,rolls=1,gsm=25,item='2TESTXXWH1N')],time_divisors=(27,{30:26,20:21})))
    assert math.isclose(r2['groups'][0]['production_hours'],91.35/27/21)
    assert '[추정] F-03: gsm 25 uses k of 20' in r2['warnings']
    assert not any('not 11 chars' in w for w in r2['warnings'])
    # Nearest choices must use only caller-supplied keys, not earlier fallback values.
    r3=verify(hp.plan([order(gsm=20),order('b',gsm=25,item='2TESTXXWH1N')]))
    assert '[추정] F-03: gsm 25 uses k of 30' in r3['warnings']
    fails(lambda: hp.plan([order()],time_divisors=(27,{})))
    print(f'  nearest gsm: 20->18 (k=21), hours={hours}; tie 25->20; item-length warning passed')


def safe_file(path):
    if not path.exists(): return "missing"
    try:
        r=subprocess.run(['xattr','-px','com.apple.metadata:_kMDItemUserTags',str(path)],capture_output=True,text=True)
        tags=plistlib.loads(bytes.fromhex(r.stdout)) if r.returncode==0 else ([] if 'No such xattr' in r.stderr else None)
        if not isinstance(tags,list) or not all(isinstance(t,str) for t in tags): return "Finder 보안 태그 판별 실패로 미검토"
        return "Finder 보안 태그로 미검토" if any(t.split('\n')[0]=='보안' for t in tags) else "safe"
    except (OSError,ValueError,plistlib.InvalidFileException): return "Finder 보안 태그 판별 실패로 미검토"


def test_optional_history():
    path=Path(__file__).resolve().parent.parent/'Process_7-8.._TEST용)_(2호기)스케줄_20260703.xlsx'
    state=safe_file(path)
    if state!='safe':
        print(f"skip historical comparison: {state}; schedule 59-row sets/length comparison NOT RUN")
        return False
    original=path.read_bytes()
    w=load_workbook(path,read_only=True,data_only=True)
    try:
        s=w['202607']
        rows=[];actual_sets=0;actual_length=0
        for r in range(7,s.max_row+1):
            width,length,rolls=[s.cell(r,c).value for c in (8,10,11)]
            if not all(isinstance(v,(int,float)) and v>0 for v in (width,length,rolls)): continue
            material,gsm,color=[s.cell(r,c).value for c in (5,6,7)]
            assert isinstance(material,str) and isinstance(gsm,(int,float)) and isinstance(color,str)
            rows.append(order(f'H-{len(rows)+1}',int(width),int(length),int(rolls),
                              item=f'2{material}{int(gsm):03d}{color}',gsm=gsm))
            mr=s.cell(r,13).value
            if mr is not None:
                assert isinstance(mr,(int,float)) and mr>0 and mr%length==0
                actual_length+=mr;actual_sets+=int(mr/length)
        assert len(rows)==59
    finally: w.close()
    result=verify(hp.plan(rows))
    sets=sum(g['sets_used'] for g in result['groups'])
    length=sum(g['total_length'] for g in result['groups'])
    assert sets<=actual_sets and length<=actual_length
    assert path.read_bytes()==original
    print(f'  history: rows=59, observed_sets={actual_sets}, planned_sets={sets}, '
          f'observed_length={actual_length}, planned_length={length}, original_unchanged=True')
    print('  history scope: cached schedule M-column comparison; separate 51-row instruction linkage NOT validated')


def make_template(path):
    w=Workbook();s=w.active;s.title='template'
    s['H8']='roll';s['I8']='kgs'
    s['D9']='WH1N';s['D10']='BK1N'
    s['H9']=99;s['J9']='OLD SYNTHETIC';s['N9']=99
    s['H11']='=SUM(H9:H10)';s['I11']='=SUM(I9:I10)'
    s.merge_cells('A11:G11')
    s['L13']='=SUM(I9:I9)';s['L14']='=I9';s['L15']='=SUM(I9)'
    w.save(path);w.close()


def test_xlsx_sets_schedule_and_formulas():
    from hanilsf_schedule import schedule
    result=verify(hp.plan([order(width=1600,rolls=3,partner_name='=1+1')],
                         constraints=[dict(type='allow_extra',order_id='a',rolls=1)]))
    timeline=schedule(result,t0='2026-01-01T00:00:00',machine_state=dict(color_code='WH1N',gsm=40,agri=False))
    with tempfile.TemporaryDirectory(prefix='.plan-test-',dir=Path(__file__).parent) as d:
        template,out=Path(d)/'template.xlsx',Path(d)/'out.xlsx'
        make_template(template);before=template.read_bytes()
        hp.to_xlsx(result,template,out,'new',schedule=timeline)
        assert template.read_bytes()==before
        w=load_workbook(out);s=w['new']
        assert w.sheetnames==['template','new','세트표','스케줄표']
        rows=[r for r in range(9,s.max_row+1) if isinstance(s.cell(r,8).value,(int,float))]
        assert sum(s.cell(r,8).value for r in rows)==4
        assert math.isclose(sum(math.prod(s.cell(r,c).value for c in (5,6,7,8))/1e6 for r in rows),result['total_kg'])
        assert s['J9'].data_type=='s' and s['J9'].value=='=1+1' and s['N9'].value is None
        assert s['S9'].value=='(1600*2)*1000 ×2세트 #세트1'
        assert s['S10'].value=='(1600*2)*1000 ×2세트 #세트1 · 초과 1롤 · 추가구매 권유'
        assert s['H12'].value=='=SUM(H9:H11)' and s['I12'].value=='=SUM(I9:I11)'
        assert s['L14'].value==s['L15'].value==s['L16'].value=='=SUM(I9:I10)'
        assert 'A12:G12' in str(s.merged_cells)
        assert w['세트표'].max_row-1==1 and w['세트표']['G2'].value==2
        assert w['스케줄표'].max_row-1==1
        assert w['스케줄표']['L2'].value==timeline['jobs'][0]['start']
        assert w['스케줄표']['J2'].value==result['groups'][0]['mr_weight_kg']
        assert w['template']['H9'].value==99
        w.close()
        fails(lambda: hp.to_xlsx(result,template,out))
        fails(lambda: hp.to_xlsx(result,template,template))
        fails(lambda: hp.to_xlsx(result,template,Path(d)/'bad.xlsx','세트표'))
        fails(lambda: hp.to_xlsx(result,template,Path(d)/'bad.xlsx',schedule=timeline|dict(complete=False)))
        # Multiple blocks retain a blank separator and expand single-cell subtotals.
        exact=hp.plan([order(width=1600,rolls=3)])
        hp.to_xlsx(exact,template,Path(d)/'exact.xlsx','exact')
        w=load_workbook(Path(d)/'exact.xlsx');s=w['exact']
        assert s['H10'].value is None and s['H9'].value+s['H11'].value==3
        assert s['L15'].value==s['L16'].value==s['L17'].value=='=SUM(I9:I11)'
        assert w['세트표'].max_row-1==2
        w.close()
    print('  xlsx: set notes/extra rows/separator/kg/formula relocation/set+schedule sheets/original preservation passed')


def test_xlsx_optional_actual_template():
    path=Path(__file__).resolve().parent.parent/'Process 3-4. PP 생산의뢰서 20260703.xlsx'
    state=safe_file(path)
    if state!='safe':
        print(f'skip actual template: {state}');return False
    original=path.read_bytes()
    result=hp.plan([order(),order('b',530)])
    with tempfile.TemporaryDirectory(prefix='.plan-test-',dir=Path(__file__).parent) as d:
        out=Path(d)/'copy.xlsx'
        hp.to_xlsx(result,path,out,'SYNTHETIC-SETS')
        w=load_workbook(out)
        assert w['세트표'].max_row-1==len(result['groups'][0]['sets'])
        s=w['SYNTHETIC-SETS']
        rows=[r for r in range(9,s.max_row+1) if isinstance(s.cell(r,8).value,(int,float))]
        assert sum(s.cell(r,8).value for r in rows)==36
        assert math.isclose(sum(math.prod(s.cell(r,c).value for c in (5,6,7,8))/1e6 for r in rows),result['total_kg'])
        w.close()
    assert path.read_bytes()==original
    print('  actual template: original unchanged, set rows/kg verified, temporary copy removed')


def test_xlsx_missing_color_block():
    from openpyxl.styles import PatternFill
    result=hp.plan([order(item='2TESTWH2N',width=1600,rolls=3),order('b',item='2TESTIV2N',rolls=1)])
    with tempfile.TemporaryDirectory(prefix='.plan-test-',dir=Path(__file__).parent) as d:
        template,out=Path(d)/'template.xlsx',Path(d)/'out.xlsx'
        make_template(template)
        w=load_workbook(template);w.active['D10'].fill=PatternFill('solid',fgColor='ABCDEF');w.save(template);w.close()
        original=template.read_bytes()
        hp.to_xlsx(result,template,out,'new')
        w=load_workbook(out);s=w['new']
        assert s['D11'].value=='WH2N' and s['D13'].value=='WH2N' and s['D14'].value=='IV2N'
        assert s['D11'].fill.fgColor.rgb=='00ABCDEF'
        assert s['H15'].value=='=SUM(H9:H14)' and s['I15'].value=='=SUM(I9:I14)'
        assert 'A15:G15' in str(s.merged_cells)
        assert sum(s.cell(r,8).value or 0 for r in range(9,15))==4
        assert 'template has no block for WH2N; appended' in result['warnings']
        assert 'template has no block for IV2N; appended' in result['warnings']
        w.close();assert template.read_bytes()==original
    print('  missing colors: WH2N/IV2N appended, last-block style/total formulas/warnings/original preservation passed')


if __name__=='__main__':
    for name,fn in list(globals().items()):
        if name.startswith('test_') and ('--core' not in sys.argv or not name.startswith('test_xlsx')):
            if fn() is not False: print('ok',name)
