"""Synthetic loader/sample tests and optional local workbook end-to-end test."""
import importlib.util
import io
import json
import tempfile
from contextlib import redirect_stdout
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from test_hanilsf_plan import safe_file, fails

assert importlib.util.find_spec('hanilsf_demo') is not None, 'demo module not implemented'
import hanilsf_demo as demo


def make_requests(path):
    w=Workbook()
    for index in range(2):
        s=w.active if index==0 else w.create_sheet()
        s.title=f'S{index+1}'
        s['H8']='roll';s['I8']='kgs'
        rows=[('WH1N',40,1000,100,3,'ASAP'),('UV4N',60,1700,200,2,datetime(2026,1,2,12)),
              ('BK1N',59,1000,100,1,'  검토 후 출고  '),('NT1N',25,500,100,2,date(2026,1,3)),
              ('WH',40,1000,100,1,None),('WH1N',40,0,100,1,None),
              ('WH1N',40,1000,100,'3',None),('WH1N',True,1000,100,1,None)]
        for r,values in enumerate(rows,9):
            color,gsm,width,length,rolls,due=values
            for c,v in {4:color,5:gsm,6:width,7:length,8:rolls,10:'SYNTHETIC',11:'TEST',12:'TEST USE',13:due}.items(): s.cell(r,c,v)
        s['H17']='=SUM(H9:H16)';s['I17']='=SUM(I9:I16)'
        s['D18']='WH1N';s['E18']=40;s['F18']=1000;s['G18']=100;s['H18']=1
    w.save(path);w.close()


def test_a_load_requests():
    with tempfile.TemporaryDirectory(prefix='.plan-test-',dir=Path(__file__).parent) as d:
        path=Path(d)/'requests.xlsx';make_requests(path);original=path.read_bytes()
        rows=demo.load_requests(path)
        assert len(rows)==8 and [r['order_id'] for r in rows]==[f'S{s}:{r}' for s in (1,2) for r in range(9,13)]
        assert rows[0]['item']=='2PD2040WH1N' and rows[0]['urgent'] is True
        assert rows[1]['item']=='2PD3060UV4N' and rows[1]['due']=='2026-01-02T12:00:00'
        assert rows[2]['item']=='2PD2059BK1N' and rows[2]['note']=='  검토 후 출고  '
        assert rows[3]['item']=='2PD2025NT1N' and rows[3]['due'].startswith('2026-01-03')
        assert rows[0]['partner_name']=='SYNTHETIC' and rows[0]['end_user']=='TEST' and rows[0]['use']=='TEST USE'
        assert len(demo.load_requests(path,sheets=['S2']))==4
        assert path.read_bytes()==original
    print('  (a) 2 sheets: 8 valid rows; D-01/date/ASAP/text conversion/filter/sheet selection/original unchanged passed')


def test_b_sample_reproducible():
    pool=[{'order_id':str(i)} for i in range(10)]
    a=demo.sample_requests(pool,4,20260915)
    assert a==demo.sample_requests(pool,4,20260915)
    assert len({r['order_id'] for r in a})==4 and len(pool)==10
    fails(lambda: demo.sample_requests(pool,11,1))
    print('  (b) seed reproducible, 4 unique samples, oversized sample rejected')


def test_main_unsolved_no_output():
    with tempfile.TemporaryDirectory(prefix='.plan-test-',dir=Path(__file__).parent) as d:
        path,out=Path(d)/'requests.xlsx',Path(d)/'out.xlsx';make_requests(path)
        log=io.StringIO()
        with redirect_stdout(log):
            code=demo.main(['--requests',str(path),'--n','2','--out',str(out),
                            '--state','WH1N,40,0','--time-limit','0.000000001'])
        assert code==1 and not out.exists() and 'UNKNOWN' in log.getvalue()
    print('  unresolved: exit=1, no XLSX written')


def test_past_due_ignore_and_keep():
    from unittest.mock import patch
    with tempfile.TemporaryDirectory(prefix='.plan-test-',dir=Path(__file__).parent) as d:
        path=Path(d)/'requests.xlsx';make_requests(path);original=path.read_bytes()
        from test_hanilsf_plan import make_template
        template=Path(d)/'template.xlsx';make_template(template)
        for mode in (None,'ignore','keep'):
            out=Path(d)/f'{mode}.xlsx';log=io.StringIO()
            argv=['--requests',str(path),'--template',str(template),'--n','8','--out',str(out),
                  '--state','WH1N,40,0','--t0','2026-01-03T00:00:00']
            if mode is not None: argv+=['--past-due',mode]
            # Spy on the real planner boundary, not a fake result: entire pipeline runs.
            with patch.object(demo,'plan',wraps=demo.plan) as planner, redirect_stdout(log):
                code=demo.main(argv)
            assert code==0 and out.exists()
            rows={r['order_id']:r for r in planner.call_args.args[0]}
            ignored=0 if mode=='keep' else 2
            summary=json.loads(log.getvalue().splitlines()[0])
            assert summary['past_due_ignored']==ignored
            for sheet in ('S1','S2'):
                assert rows[f'{sheet}:9']['urgent'] is True
                assert ('due' in rows[f'{sheet}:10'])==(mode=='keep')
                if mode=='keep': assert rows[f'{sheet}:10']['due']=='2026-01-02T12:00:00'
                assert rows[f'{sheet}:12']['due']=='2026-01-03T00:00:00'  # equality is not past
            assert path.read_bytes()==original
        assert 'due' in demo.load_requests(path)[1]  # loader remains lossless
    print('  past due: default/ignore removed=2, keep removed=0; equal-t0 due/urgent/original preserved')


def test_missing_requests_error_details():
    with tempfile.TemporaryDirectory(prefix='.plan-test-',dir=Path(__file__).parent) as d:
        missing,out=Path(d)/'missing.xlsx',Path(d)/'out.xlsx';log=io.StringIO()
        with redirect_stdout(log):
            code=demo.main(['--requests',str(missing),'--out',str(out),'--state','WH1N,40,0'])
        message=json.loads(log.getvalue())
        assert code==1 and not out.exists() and message['output'] is None
        assert message['error'].startswith('FileNotFoundError: ') and 'missing.xlsx' in message['error']
    print('  missing requests: exit=1, error includes FileNotFoundError and message, output=null')


def test_c_optional_end_to_end():
    path=Path(__file__).resolve().parent.parent/'Process 3-4. PP 생산의뢰서 20260703.xlsx'
    state=safe_file(path)
    if state!='safe':
        print(f'skip optional end-to-end: {state}');return False
    original=path.read_bytes()
    with tempfile.TemporaryDirectory(prefix='.plan-test-',dir=Path(__file__).parent) as d:
        out=Path(d)/'end-to-end.xlsx'
        log=io.StringIO()
        with redirect_stdout(log):
            code=demo.main(['--requests',str(path),'--n','30','--seed','20260915','--out',str(out),
                            '--t0','2026-09-16T00:00:00','--state','WH1N,40,0','--time-limit','30'])
        # Only summaries are printed; no customer or worksheet-row details.
        print(log.getvalue(),end='')
        messages=[json.loads(line) for line in log.getvalue().splitlines()]
        population=messages[0];stages={m['stage']:m for m in messages if 'stage' in m}
        if (population['pool_rows'],population['sample_items'],stages.get(2,{}).get('jobs'))!=(412,16,31):
            print('  reference differs: inspect workbook revision/valid-row population or equal-optimum pattern selection')
        assert code==0
        w=load_workbook(out)
        assert '세트표' in w.sheetnames and '스케줄표' in w.sheetnames and len(w.sheetnames)==22
        assert w['스케줄표'].max_row>1
        print(f"  (c) end-to-end: exit={code}, added_sheets=3, schedule_rows={w['스케줄표'].max_row-1}")
        w.close()
    assert path.read_bytes()==original
    print('  (c) original bytes unchanged; temporary output removed')


if __name__=='__main__':
    for name,fn in list(globals().items()):
        if name.startswith('test_') and fn() is not False: print('ok',name)
